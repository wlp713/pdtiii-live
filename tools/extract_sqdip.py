#!/usr/bin/env python3
"""Extract the current monthly SQDIP workbook totals into a static browser data file."""
from __future__ import annotations
import argparse
import json
import re
from pathlib import Path
from datetime import date, datetime

SHEETS = {
    "PRO.1": ("p1.xls", "SQDIP"),
    "PRO.2": ("p2.xlsx", "SQDIP"),
    "PRO.3": ("p3.xlsx", "SQDIP"),
    "PRO.4-5": ("p45.xlsx", "SQDIP (ข้อมูล)"),
    "PRO.6": ("p6.xlsx", "SQDIP"),
}
DIMS = ("S", "Q", "D", "I", "P")


def clean_value(value):
    if value is None:
        return None
    if isinstance(value, str):
        s = value.strip()
        if not s or s.startswith("#"):
            return None
        try:
            return float(s.replace(",", ""))
        except ValueError:
            return s
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def open_sheet(path: Path, sheet_name: str):
    if path.suffix.lower() == ".xls":
        import xlrd
        book = xlrd.open_workbook(str(path), on_demand=True)
        sheet = book.sheet_by_name(sheet_name)
        def get(row, col):
            if row >= sheet.nrows or col >= sheet.ncols:
                return None
            return sheet.cell_value(row, col)
        row_count = sheet.nrows
        col_count = sheet.ncols
    else:
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        sheet = book[sheet_name]
        def get(row, col):
            return sheet.cell(row + 1, col + 1).value
        row_count = sheet.max_row
        col_count = sheet.max_column
    total_col = None
    for r in range(min(row_count, 20)):
        for c in range(col_count):
            value = get(r, c)
            if isinstance(value, str) and value.strip().casefold() == "total":
                total_col = c
                break
        if total_col is not None:
            break
    if total_col is None:
        raise ValueError(f"Cannot find Total column in {path.name}/{sheet_name}")
    return book, get, row_count, col_count, total_col


def dimension_of(value):
    if not isinstance(value, str):
        return None
    match = re.match(r"^\s*([SQDIP])(?=\s|$|\n)", value.upper())
    return match.group(1) if match else None


def parse_number(text, pattern):
    match = re.search(pattern, text, re.I)
    if not match:
        return None
    return float(match.group(1).replace(",", "")), (match.group(2) if match.lastindex and match.lastindex >= 2 else "").strip()


def title_lines(header):
    lines = [part.strip() for part in str(header or "").splitlines() if part.strip()]
    if lines and re.match(r"^[SQDIP](?:\s|$)", lines[0], re.I):
        lines[0] = re.sub(r"^[SQDIP]\s*", "", lines[0], flags=re.I).strip()
    return [line for line in lines if line and not re.search(r"\b(?:target|cost)\b|目标", line, re.I)]


def short_title(shop, dim, header):
    raw = str(header or "")
    if shop == "PRO.6" and dim == "Q":
        match = re.search(r"NG\s+([A-Za-z-]+)", raw, re.I)
        if match:
            return "NG " + match.group(1)
    lines = title_lines(header)
    joined = " ".join(lines)
    if shop == "PRO.4-5" and dim == "Q" and "Production 4" in joined:
        return "Production 4"
    if dim == "Q" and "ng line out shipping" in joined.casefold():
        return "NG Line Out Shipping"
    if dim == "Q" and ("ng inprocess" in joined.casefold() or "ng inprocess" in joined.casefold()):
        return "NG Inprocess"
    if dim == "Q" and "งานเชื่อมเสีย" in joined:
        return "焊接不良率"
    if shop == "PRO.6" and dim == "Q":
        match = re.search(r"NG\s+([A-Za-z-]+)", joined, re.I)
        if match:
            return "NG " + match.group(1)
    if dim == "P":
        return "UPPH"
    candidates = []
    for line in lines:
        if line.upper() in {"S-SAFETY", "5S IMPROVEMENT", "NG INPROCESS", "NG LINE OUT SHIPPING", "LOSS", "EFFICIENCY"}:
            if line.upper() in {"NG INPROCESS", "NG LINE OUT SHIPPING"}:
                candidates.append(line)
            continue
        if re.search(r"(闭环率|安全事故|作业完成率|品质性能不良率|品质焊接不良率|损耗|效率|ประสิทธิภาพ|S-Safety)", line, re.I):
            candidates.append(line)
        elif len(line) <= 38:
            candidates.append(line)
    if candidates:
        value = candidates[0]
        if "งานเชื่อมเสีย" in value:
            return "焊接不良率"
        return value[:38]
    return (lines[0] if lines else dim)[:38]


def row_entry(label, value):
    label = str(label or "").strip()
    if not label:
        return None
    cleaned = clean_value(value)
    return {"label": label, "value": cleaned, "sourceError": isinstance(value, str) and value.strip().startswith("#")}


def choose_actual(shop, dim, header, rows):
    text = (str(header or "") + " " + " ".join(r["label"] for r in rows)).casefold()
    labels = [(r, r["label"].casefold()) for r in rows]
    def find(terms):
        for term in terms:
            for row, label in labels:
                if term.casefold() in label:
                    return row
        return None
    if dim == "S":
        if "s-safety" in text or "安全事故" in text:
            return find(("安全事故", "zero"))
        if "target10list" in text or "target 10 list" in text:
            return find(("项目数", "list"))
        return find(("闭环率",)) or find(("闭环数", "项目数"))
    if dim == "Q":
        if shop == "PRO.4-5":
            return find(("cost (thb)", "不良数"))
        ppm = find(("ppm",))
        if ppm:
            return ppm
        pct = find(("%", "率",))
        if pct:
            return pct
        return find(("不良数", "defect", "cost"))
    if dim == "D":
        if find(("target",)) and find(("完成数",)):
            return find(("完成数",))
        if "loss pro.1" in text:
            return find(("损耗率", "loss pro.2"))
        if "loss ng pro.2" in text:
            return find(("loss pro.2", "完成率"))
        return find(("完成率", "损耗率", "loss pro.2", "损失"))
    if dim == "I":
        if "energy cost" in text:
            return find(("损耗率", "不良金额", "完成率"))
        if "loss cost" in text:
            return find(("不良金额", "损耗率", "完成率"))
        return find(("完成率", "损耗率", "不良金额", "wait parts"))
    if dim == "P":
        return next((r for r, label in labels if "upph" in label and not label.strip().startswith("target")), None)
    return None


def format_kind(label, dim, header, actual=None):
    text = (str(label or "") + " " + str(header or "")).casefold()
    label_text = str(label or "").casefold()
    if "upph" in text:
        return "upph"
    if "ppm" in label_text:
        return "ppm"
    if "thb" in text or "cost (" in label_text:
        return "money"
    if any(word in label_text for word in ("闭环率", "完成率")):
        return "ratio-percent"
    if "%" in label_text:
        return "percent-point"
    if "zero" in label_text or "安全事故" in label_text or "项目数" in label_text or "完成数" in label_text or "不良数" in label_text:
        return "count"
    return "number"


def target_for(shop, dim, header, rows, actual, helpers):
    labels = [(r, r["label"].casefold()) for r in rows]
    # An explicit target inside the same source block has priority.
    target_row = next((r for r, label in labels if re.fullmatch(r"target(?:\s+upph)?", label.strip())), None)
    if target_row:
        return {"value": target_row["value"], "label": target_row["label"], "format": format_kind(actual["label"], dim, header)}
    metric = actual["label"].casefold()
    title = str(header or "")
    low = title.casefold()
    if dim == "S" and ("s-safety" in low or "安全事故" in low):
        if re.search(r"\bzero\b", low) or "zero" in actual["label"].casefold() or any("zero" in r["label"].casefold() for r in rows):
            return {"value": 0, "label": "Zero", "format": "count"}
    if dim == "S" and ("5s" in low or "improvement" in low or "改善" in low):
        val = parse_number(title, r"target\s*([\d,.]+)\s*list\s*/\s*month")
        if val:
            return {"value": val[0], "label": "Target / month", "format": "count", "unit": "项/月"}
        val = parse_number(title, r"([\d,.]+)\s*%\s*/\s*month")
        if val:
            return {"value": val[0], "label": "Target / month", "format": "percent-point", "unit": "%/月"}
    # Targets repeated in a source row whose label includes the KPI name.
    if "upph" in metric:
        for r in rows:
            if "target" in r["label"].casefold() and "upph" in r["label"].casefold():
                return {"value": r["value"], "label": r["label"], "format": "upph"}
        val = parse_number(actual["label"], r"UPPH\s+([\d,.]+)")
        if not val:
            val = parse_number(title, r"UPPH\s+([\d,.]+)")
        if val:
            return {"value": val[0], "label": "UPPH target", "format": "upph"}
        for r in helpers:
            if "upph target" in r["label"].casefold():
                return {"value": r["value"], "label": r["label"], "format": "upph"}
    # Header targets retain the source's stated unit and are only used when present.
    val = parse_number(title, r"target\s*[:：]?\s*([\d,.]+)\s*(%|THB(?:\s*/\s*month)?|set(?:\s*/\s*day)?|list\s*/\s*month)?")
    if val:
        unit = val[1].replace(" ", "")
        fmt = "percent-point" if "%" in unit else ("money" if "thb" in unit.casefold() else format_kind(actual["label"], dim, title))
        return {"value": val[0], "label": "Target", "format": fmt, "unit": unit}
    # P3 keeps the target lines outside the metric blocks in the same sheet.
    for r in helpers:
        label = r["label"].casefold()
        if dim == "D" and "energy target" in label and "energy" in low:
            return {"value": r["value"], "label": r["label"], "format": "percent-point", "unit": "%"}
        if dim == "Q" and "q " in label and "target" in label and ("weld" in low or "งานเชื่อม" in low):
            return {"value": r["value"], "label": r["label"], "format": "percent-point", "unit": "%"}
        if dim == "S" and "5s improvement" in label and "5s" in low and not re.search(r"\d+\s*%\s*/\s*month", low):
            return {"value": r["value"], "label": r["label"], "format": "count"}
    return None


def read_workbook(shop, input_dir):
    file_name, sheet_name = SHEETS[shop]
    path = input_dir / file_name
    if not path.exists():
        raise FileNotFoundError(f"Missing source workbook: {path}")
    book, get, row_count, col_count, total_col = open_sheet(path, sheet_name)
    groups = []
    helpers = []
    current = None
    blank_run = 0
    for row in range(min(row_count, 300)):
        b = get(row, 1)
        c = get(row, 2)
        if not (str(b or "").strip() or str(c or "").strip()):
            blank_run += 1
            if blank_run >= 2:
                current = None
            continue
        blank_run = 0
        if isinstance(c, str) and re.search(r"target|目标", c, re.I):
            helper = row_entry(c, get(row, total_col))
            if helper:
                helpers.append(helper)
        dim = dimension_of(b)
        if dim:
            current = {"dimension": dim, "header": b, "rows": []}
            groups.append(current)
            entry = row_entry(c, get(row, total_col))
            if entry and not re.match(r"^类别/目标", entry["label"], re.I):
                current["rows"].append(entry)
            continue
        if current:
            entry = row_entry(c, get(row, total_col))
            if entry and not re.match(r"^类别/目标", entry["label"], re.I):
                current["rows"].append(entry)
    for group in groups:
        dim = group["dimension"]
        if shop == "PRO.4-5" and dim == "Q" and "production 4" not in str(group["header"]).casefold():
            continue
        actual = choose_actual(shop, dim, group["header"], group["rows"])
        if not actual:
            continue
        actual["format"] = format_kind(actual["label"], dim, group["header"])
        target = target_for(shop, dim, group["header"], group["rows"], actual, helpers)
        components = [r for r in group["rows"] if not re.match(r"^target\b", r["label"].strip(), re.I)]
        group.update({
            "title": short_title(shop, dim, group["header"]),
            "metricLabel": actual["label"],
            "actual": actual,
            "target": target,
            "components": components,
        })
        del group["rows"]
    return book, [g for g in groups if "actual" in g]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    output = {"period": "2026-10", "plants": []}
    for shop in SHEETS:
        book, groups = read_workbook(shop, args.input_dir)
        grouped = {dim: [] for dim in DIMS}
        for group in groups:
            grouped[group["dimension"]].append(group)
        output["plants"].append({"id": shop, "source": SHEETS[shop][0], "groups": grouped})
        if hasattr(book, "release_resources"):
            book.release_resources()
        elif hasattr(book, "close"):
            book.close()
    payload = json.dumps(output, ensure_ascii=False, separators=(",", ":"), allow_nan=False)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        "/* Generated from the current SQDIP monthly workbooks; do not edit by hand. */\n"
        "(function(){var data=" + payload + ";window.dispatchEvent(new CustomEvent('sqdip:data-ready',{detail:data}));})();\n",
        encoding="utf-8",
    )
    print(f"Wrote {args.output} ({len(payload):,} JSON characters)")
    for plant in output["plants"]:
        summary = ", ".join(f"{dim}:{len(plant['groups'][dim])}" for dim in DIMS)
        print(f"{plant['id']} {summary}")


if __name__ == "__main__":
    main()